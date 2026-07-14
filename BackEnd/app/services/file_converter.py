import pandas as pd
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter
from io import BytesIO
from typing import List, Dict, Any
import logging
from app.utils.header_utils import normalize_header

logger = logging.getLogger(__name__)

class FileConverterService:
    @staticmethod
    def _is_blank(value: Any) -> bool:
        if value is None:
            return True
        if isinstance(value, str):
            return not value.strip()
        return False

    @staticmethod
    def _get_merged_cell_value(ws, row, col):
        """
        Checks if a cell is part of a merged range and returns the
        value of the top-left cell of that range.
        """
        for merged_range in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged_range.bounds
            if min_row <= row <= max_row and min_col <= col <= max_col:
                return ws.cell(row=min_row, column=min_col).value
        return ws.cell(row=row, column=col).value

    @staticmethod
    def _row_signature(ws, row_idx):
        row_values = [FileConverterService._get_merged_cell_value(ws, row_idx, col) for col in range(1, ws.max_column + 1)]
        text_values = [str(v).strip() for v in row_values if not FileConverterService._is_blank(v)]
        non_null_count = len(text_values)
        if non_null_count == 0:
            return {
                "row_idx": row_idx,
                "row_values": row_values,
                "text_values": [],
                "non_null_count": 0,
                "avg_len": 0,
                "uppercase_like_count": 0,
                "keyword_hits": 0,
                "score": 0,
            }

        avg_len = sum(len(v) for v in text_values) / non_null_count
        uppercase_like_count = sum(1 for v in text_values if any(ch.isalpha() for ch in v) and v == v.upper())
        keyword_hits = sum(
            1
            for v in text_values
            if any(keyword in v.upper() for keyword in [
                "CÓDIGO", "CODIGO", "NÚMERO", "NUMERO", "DESCRIPCIÓN", "DESCRIPCION",
                "MARCA", "ESTADO", "COSTO", "FECHA", "NOMBRE", "FOLIO", "PROPIEDAD",
                "RECURSO", "SERIE", "PROVEEDOR", "GASTO", "OBJETO", "DIRECCIÓN", "DIRECCION"
            ])
        )
        short_text_count = sum(1 for v in text_values if len(v) <= 40)
        score = (
            non_null_count * 1.4
            + uppercase_like_count * 0.6
            + keyword_hits * 2.0
            + short_text_count * 0.15
            - max(avg_len - 28, 0) * 0.35
        )

        return {
            "row_idx": row_idx,
            "row_values": row_values,
            "text_values": text_values,
            "non_null_count": non_null_count,
            "avg_len": avg_len,
            "uppercase_like_count": uppercase_like_count,
            "keyword_hits": keyword_hits,
            "score": score,
        }

    @staticmethod
    def _detect_header_rows(ws):
        """
        Detects a multi-row header block dynamically.
        """
        candidates = [FileConverterService._row_signature(ws, row_idx) for row_idx in range(1, min(121, ws.max_row + 1))]
        non_empty_candidates = [candidate for candidate in candidates if candidate["non_null_count"] > 0]
        if not non_empty_candidates:
            return [1]

        primary = max(non_empty_candidates, key=lambda candidate: candidate["score"])
        min_score = max(8.0, primary["score"] * 0.45)

        header_rows = [primary["row_idx"]]

        current_row = primary["row_idx"] - 1
        while current_row >= 1:
            candidate = next((item for item in candidates if item["row_idx"] == current_row), None)
            if not candidate or candidate["non_null_count"] == 0:
                break
            if candidate["score"] < min_score:
                break
            header_rows.insert(0, current_row)
            current_row -= 1

        current_row = primary["row_idx"] + 1
        while current_row <= min(ws.max_row, 120):
            candidate = next((item for item in candidates if item["row_idx"] == current_row), None)
            if not candidate or candidate["non_null_count"] == 0:
                break
            if candidate["score"] < min_score:
                break
            header_rows.append(current_row)
            current_row += 1

        return sorted(set(header_rows))

    @staticmethod
    def _build_row_cells(ws, row_idx):
        cells = []
        values = []

        for col_idx in range(1, ws.max_column + 1):
            value = FileConverterService._get_merged_cell_value(ws, row_idx, col_idx)
            if not FileConverterService._is_blank(value):
                values.append(value)
            cells.append({
                "column_index": col_idx,
                "column_letter": get_column_letter(col_idx),
                "value": value,
            })

        return cells, values

    @staticmethod
    def _combine_header_rows(header_rows, ws, col_idx):
        original_parts = []
        normalized_parts = []

        for header_row_idx in header_rows:
            header_value = FileConverterService._get_merged_cell_value(ws, header_row_idx, col_idx)
            if FileConverterService._is_blank(header_value):
                continue
            header_text = str(header_value).strip()
            original_parts.append(header_text)
            normalized_parts.append(normalize_header(header_text))

        if not normalized_parts:
            return "unnamed_column", "unnamed"

        normalized_header = "__".join(normalized_parts)
        original_header = " | ".join(original_parts)
        return normalized_header, original_header

    @staticmethod
    async def convert_to_json(file_content: bytes, filename: str) -> List[Dict[str, Any]]:
        """
        Converts CSV or Excel file content to a segmented JSON structure.
        Segments data into: document_metadata_raw, table_rows_raw, summary_rows_raw, and footer_metadata_raw.
        """
        try:
            all_sheets_processed = []
            if filename.endswith('.csv'):
                df = pd.read_csv(BytesIO(file_content))
                cleaned_df = df.where(pd.notnull(df), None)
                all_sheets_processed.append({
                    "sheet_name": "default",
                    "document_metadata_raw": [],
                    "table_rows_raw": cleaned_df.to_dict(orient='records'),
                    "summary_rows_raw": [],
                    "footer_metadata_raw": [],
                    "form_selections_raw": []
                })
            
            elif filename.endswith(('.xlsx', '.xls')):
                wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
                
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    
                    # 1. Detect Header Block
                    header_row_indices = FileConverterService._detect_header_rows(ws)
                    header_row_idx = header_row_indices[0]
                    last_header_row_idx = header_row_indices[-1]
                    
                    # 2. Extract Document Metadata (Rows before header)
                    metadata_raw = []
                    sheet_cells_raw = []
                    for row_idx in range(1, ws.max_row + 1):
                        row_cells, row_values = FileConverterService._build_row_cells(ws, row_idx)
                        non_empty_count = sum(1 for value in row_values if not FileConverterService._is_blank(value))
                        if non_empty_count == 0:
                            continue

                        row_role = "body"
                        if row_idx < header_row_idx:
                            row_role = "document_metadata"
                        elif row_idx in header_row_indices:
                            row_role = "header"
                        elif row_idx > last_header_row_idx and any(val and any(keyword in str(val).upper() for keyword in ["TOTAL", "SUBTOTAL", "SUMA", "GRAN TOTAL"]) for val in row_values):
                            row_role = "summary"

                        sheet_cells_raw.append({
                            "_row_index": row_idx,
                            "_row_role_guess": row_role,
                            "_non_empty_cell_count": non_empty_count,
                            "cells": row_cells,
                        })

                        if row_idx < header_row_idx:
                            row_data = {
                                f"col_{col}": FileConverterService._get_merged_cell_value(ws, row_idx, col)
                                for col in range(1, ws.max_column + 1)
                            }
                            metadata_raw.append({
                                "_row_index": row_idx,
                                "_row_role_guess": row_role,
                                "_non_empty_cell_count": non_empty_count,
                                "data": row_data,
                            })
                    
                    # 3. Extract Headers for the table
                    raw_headers = [
                        FileConverterService._combine_header_rows(header_row_indices, ws, col)
                        for col in range(1, ws.max_column + 1)
                    ]
                    
                    # 4. Extract Rows and Classify (Table vs Summary vs Footer)
                    table_rows = []
                    summary_rows = []
                    footer_rows = []
                    
                    for row_idx in range(last_header_row_idx + 1, ws.max_row + 1):
                        row_dict = {}
                        row_values = []
                        row_cells = []
                        
                        for col_idx, (normalized_header, original_header) in enumerate(raw_headers, start=1):
                            val = FileConverterService._get_merged_cell_value(ws, row_idx, col_idx)
                            row_values.append(val)
                            row_cells.append({
                                "column_index": col_idx,
                                "column_letter": get_column_letter(col_idx),
                                "value": val,
                                "header_normalized": normalized_header,
                                "header_original": original_header,
                            })
                            
                            row_dict[normalized_header] = {
                                "value": val,
                                "original_header": original_header if original_header else "unnamed"
                            }
                        
                        # Exhaustive Classification Logic
                        # a) Summary Row
                        if any(val and any(keyword in str(val).upper() for keyword in ["TOTAL", "SUBTOTAL", "SUMA", "GRAN TOTAL"]) for val in row_values):
                            summary_rows.append({"_row_index": row_idx, "_row_role_guess": "summary", "cells": row_cells, "data": row_dict})
                        
                        # b) Table Record: Any row with at least one value is a potential record
                        elif any(v is not None for v in row_values):
                            # If it's a table record, add it to table_rows
                            if any(v['value'] is not None for v in row_dict.values()):
                                table_rows.append({
                                    "_row_index": row_idx,
                                    "_row_role_guess": "body",
                                    "_non_empty_cell_count": sum(1 for v in row_values if not FileConverterService._is_blank(v)),
                                    "cells": row_cells,
                                    "data": row_dict,
                                })
                            
                            # Additionally, if the row is sparse (looks like a footer/comment), 
                            # capture it as raw footer data too for AI safety
                            non_null_count = sum(1 for v in row_values if v is not None)
                            if non_null_count < (len(raw_headers) // 2):
                                footer_data = {f"col_{col}": FileConverterService._get_merged_cell_value(ws, row_idx, col) 
                                              for col in range(1, ws.max_column + 1)}
                                footer_rows.append({
                                    "_row_index": row_idx,
                                    "_row_role_guess": "footer",
                                    "cells": row_cells,
                                    "data": footer_data,
                                })
                        
                        # c) Fallback: Row exists but didn't fit table/summary logic
                        else:
                            footer_data = {f"col_{col}": FileConverterService._get_merged_cell_value(ws, row_idx, col) 
                                          for col in range(1, ws.max_column + 1)}
                            if any(v is not None for v in footer_data.values()):
                                footer_rows.append({
                                    "_row_index": row_idx,
                                    "_row_role_guess": "footer",
                                    "cells": row_cells,
                                    "data": footer_data,
                                })
                    
                    all_sheets_processed.append({
                        "sheet_name": sheet_name,
                        "detected_header_rows": header_row_indices,
                        "detected_data_start_row": last_header_row_idx + 1,
                        "sheet_cells_raw": sheet_cells_raw,
                        "document_metadata_raw": metadata_raw,
                        "table_rows_raw": table_rows,
                        "summary_rows_raw": summary_rows,
                        "footer_metadata_raw": footer_rows,
                        "form_selections_raw": []
                    })
            else:
                raise ValueError("Unsupported file format. Please provide a .csv, .xlsx, or .xls file.")

            return all_sheets_processed
        except Exception as e:
            logger.error(f"Error procesando el archivo: {str(e)}")
            raise e
